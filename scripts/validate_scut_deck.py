#!/usr/bin/env python
from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


EMU_PER_INCH = 914400
EXPECTED_SIZE = (13.33, 7.5)
EXPECTED_CHAPTERS = ["研究背景", "问题与贡献", "方法框架", "实验设计", "结果讨论", "结论展望"]
PLACEHOLDERS = ["鲤工仔", "Your Title Here SCUT", "XX年XX月", "华工蓝幻灯片模板", "待补充论文要点"]


def iter_shapes(shapes: Any):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def shape_text(shape: Any) -> str:
    return getattr(shape, "text", "") or ""


def all_text(prs: Presentation) -> str:
    return "\n".join(shape_text(shape) for slide in prs.slides for shape in iter_shapes(slide.shapes))


def load_plan(path: Path | None, pptx: Path) -> dict[str, Any] | None:
    if path is None:
        candidate = pptx.with_suffix(".slide_plan.json")
        path = candidate if candidate.exists() else None
    if path is None:
        return None
    return json.loads(path.read_text(encoding="utf-8"))


def check(condition: bool, message: str, errors: list[str], ok: list[str]) -> None:
    if condition:
        ok.append(message)
    else:
        errors.append(message)


def validate(pptx: Path, plan_path: Path | None = None) -> tuple[list[str], list[str]]:
    prs = Presentation(str(pptx))
    plan = load_plan(plan_path, pptx)
    text = all_text(prs)
    errors: list[str] = []
    ok: list[str] = []

    width = round(prs.slide_width / EMU_PER_INCH, 2)
    height = round(prs.slide_height / EMU_PER_INCH, 2)
    check((width, height) == EXPECTED_SIZE, f"canvas is {EXPECTED_SIZE[0]} x {EXPECTED_SIZE[1]} in", errors, ok)
    check(len(prs.slides) >= 14, "deck has at least 14 slides for six chapters", errors, ok)

    for placeholder in PLACEHOLDERS:
        check(placeholder not in text, f"no placeholder remains: {placeholder}", errors, ok)

    for idx, chapter in enumerate(EXPECTED_CHAPTERS, 1):
        check(chapter in text, f"chapter title present: {idx:02d} {chapter}", errors, ok)
        check(f"{idx:02d}" in text, f"chapter number present: {idx:02d}", errors, ok)

    if plan is not None:
        slides = plan.get("slides", [])
        sections = [slide for slide in slides if slide.get("slide_type") == "section"]
        toc = next((slide for slide in slides if slide.get("slide_type") == "toc"), None)
        check(len(slides) == len(prs.slides), "slide_plan count matches PPT slide count", errors, ok)
        check(len(sections) == 6, "slide_plan has six section dividers", errors, ok)
        check([section.get("section_no") for section in sections] == list(range(1, 7)), "section numbers are 1-6", errors, ok)
        check(toc is not None and toc.get("body") == EXPECTED_CHAPTERS, "TOC body lists the six expected chapters", errors, ok)

    return errors, ok


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Validate a generated SCUT paper PPT deck.")
    parser.add_argument("pptx", type=Path, help="Generated PPTX path.")
    parser.add_argument("--plan", type=Path, help="Optional slide_plan.json path. Defaults to sibling file.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    errors, ok = validate(args.pptx, args.plan)
    for message in ok:
        print(f"OK  {message}")
    for message in errors:
        print(f"ERR {message}", file=sys.stderr)
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
