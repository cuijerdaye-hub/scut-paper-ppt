#!/usr/bin/env python
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import subprocess
import sys
from copy import deepcopy
from pathlib import Path
from typing import Any

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt


SCRIPT_DIR = Path(__file__).resolve().parent
SKILL_DIR = SCRIPT_DIR.parent
DEFAULT_TEMPLATE = SKILL_DIR / "assets" / "scut-blue-template.pptx"
WHITE = (255, 255, 255)
BLACK = (0, 0, 0)
SCUT_BLUE = (45, 79, 143)
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

SLOGANS = {
    "厚德尚学 自强不息",
    "务实创新 追求卓越",
    "厚德尚学 自强不息 务实创新 追求卓越",
}

PLACEHOLDER_TOKENS = [
    "鲤工仔",
    "Your Title Here SCUT",
    "XX年XX月",
    "华工蓝幻灯片模板",
]

CATEGORY_KEYWORDS = {
    "background": ["背景", "引言", "绪论", "现状", "相关工作", "introduction", "background", "related work"],
    "problem": ["问题", "挑战", "动机", "贡献", "创新", "不足", "challenge", "problem", "contribution"],
    "method": ["方法", "模型", "框架", "算法", "设计", "method", "model", "framework", "approach"],
    "experiment": ["实验", "数据", "评估", "验证", "设置", "experiment", "dataset", "evaluation"],
    "result": ["结果", "发现", "性能", "实验结果", "result", "finding", "performance"],
    "discussion": ["讨论", "局限", "启示", "消融", "discussion", "limitation", "ablation"],
    "conclusion": ["结论", "总结", "展望", "未来", "conclusion", "future"],
}

def clip(text: str, limit: int) -> str:
    text = normalize_space(text)
    if len(text) <= limit:
        return text
    return text[: max(1, limit - 1)].rstrip("，,；;。.") + "…"


def normalize_space(text: str) -> str:
    return re.sub(r"\s+", " ", text or "").strip()


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return read_pdf_text(path)
    if suffix in {".md", ".markdown", ".txt"}:
        for enc in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return path.read_text(encoding=enc)
            except UnicodeDecodeError:
                continue
        return path.read_text(errors="ignore")
    raise ValueError(f"Unsupported input format: {path.suffix}. Use PDF, Markdown, or TXT.")


def read_pdf_text(path: Path) -> str:
    try:
        from markitdown import MarkItDown

        result = MarkItDown().convert(str(path))
        content = getattr(result, "text_content", None)
        if content:
            return content
    except Exception:
        pass

    completed = subprocess.run(
        [sys.executable, "-m", "markitdown", str(path)],
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
        errors="replace",
    )
    return completed.stdout


def strip_markdown(line: str) -> str:
    line = re.sub(r"!\[[^\]]*\]\([^)]+\)", "", line)
    line = re.sub(r"\[[^\]]+\]\([^)]+\)", "", line)
    line = re.sub(r"[*_`>#\-]+", " ", line)
    return normalize_space(line)


def clean_source_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_title(text: str) -> str:
    match = re.search(r"^\s*#\s+(.+?)\s*$", text, flags=re.MULTILINE)
    if match:
        return clip(strip_markdown(match.group(1)), 36)

    for line in text.splitlines():
        candidate = strip_markdown(line)
        if not candidate:
            continue
        if re.match(r"^(摘要|关键词|abstract|keywords|作者|author)\b", candidate, flags=re.I):
            continue
        if 4 <= len(candidate) <= 80:
            return clip(candidate, 36)
    return "论文研究汇报"


def extract_named_block(text: str, names: list[str], limit: int = 420) -> str:
    lines = text.splitlines()
    start = None
    inline = ""
    name_pat = "|".join(re.escape(name) for name in names)
    for idx, line in enumerate(lines):
        stripped = strip_markdown(line)
        if re.match(rf"^({name_pat})\s*[:：]\s*(.+)$", stripped, flags=re.I):
            inline = re.sub(rf"^({name_pat})\s*[:：]\s*", "", stripped, flags=re.I)
            start = idx + 1
            break
        if re.match(rf"^(#+\s*)?({name_pat})\s*$", stripped, flags=re.I):
            start = idx + 1
            break
    if start is None:
        return ""

    collected = [inline] if inline else []
    for line in lines[start:]:
        stripped = strip_markdown(line)
        if not stripped:
            if collected:
                break
            continue
        if is_heading(line) or re.match(r"^(关键词|keywords)\s*[:：]?", stripped, flags=re.I):
            break
        collected.append(stripped)
        if sum(len(x) for x in collected) > limit:
            break
    return clip(" ".join(collected), limit)


def is_heading(line: str) -> str | None:
    raw = line.strip()
    if not raw:
        return None
    md = re.match(r"^#{1,4}\s+(.{2,80})$", raw)
    if md:
        return strip_markdown(md.group(1))
    if len(raw) > 70 or raw.endswith(("。", "；", ";")):
        return None
    numbered = re.match(
        r"^(?:第?[一二三四五六七八九十]+[章节部分]?[\s、：:.]|[一二三四五六七八九十]+、|\d+(?:\.\d+)*[\s、.]+)(.{2,60})$",
        raw,
    )
    if numbered:
        return strip_markdown(numbered.group(1))
    return None


def parse_sections(text: str) -> list[dict[str, str]]:
    sections: list[dict[str, str]] = []
    title = "正文"
    buf: list[str] = []
    for line in text.splitlines():
        heading = is_heading(line)
        if heading:
            if buf:
                sections.append({"title": title, "text": "\n".join(buf).strip()})
            title = heading
            buf = []
        else:
            cleaned = strip_markdown(line)
            if cleaned:
                buf.append(cleaned)
    if buf:
        sections.append({"title": title, "text": "\n".join(buf).strip()})
    return sections or [{"title": "正文", "text": strip_markdown(text)}]


def classify_sections(sections: list[dict[str, str]], all_text: str) -> dict[str, str]:
    buckets: dict[str, list[str]] = {key: [] for key in CATEGORY_KEYWORDS}

    # Prefer explicit section headings. Abstract text often names methods and
    # results, but it should not swallow the body slides.
    for section in sections:
        title = section["title"].lower()
        for category, keywords in CATEGORY_KEYWORDS.items():
            if any(keyword.lower() in title for keyword in keywords):
                buckets[category].append(f"{section['title']}。{section['text']}")

    for category, parts in buckets.items():
        if parts:
            continue
        for section in sections:
            title = section["title"].lower()
            if any(skip in title for skip in ["摘要", "abstract", "关键词", "keywords"]):
                continue
            body = section["text"][:500].lower()
            if any(keyword.lower() in body for keyword in CATEGORY_KEYWORDS[category]):
                parts.append(f"{section['title']}。{section['text']}")
                break

    fallback = all_text
    return {
        category: "\n".join(parts).strip() if parts else fallback
        for category, parts in buckets.items()
    }


def split_sentences(text: str) -> list[str]:
    text = re.sub(r"\[[0-9,\-\s]+\]", "", text)
    primary_parts = re.split(r"[。！？!?；;\n]+|(?<=[.])\s+", text)
    parts: list[str] = []
    for part in primary_parts:
        if len(part) > 80:
            parts.extend(re.split(r"[，,]", part))
        else:
            parts.append(part)
    seen: set[str] = set()
    sentences: list[str] = []
    for part in parts:
        item = strip_markdown(part)
        if len(item) < 8:
            continue
        item = clip(item, 90)
        key = item.lower()
        if key not in seen:
            seen.add(key)
            sentences.append(item)
    return sentences


def bullets_from(text: str, count: int, max_len: int = 48) -> list[str]:
    bullet_lines: list[str] = []
    for line in text.splitlines():
        if re.match(r"^\s*([-*+]|\d+[.)、])\s+", line):
            bullet_lines.append(strip_markdown(line))
    candidates = bullet_lines + split_sentences(text)
    result: list[str] = []
    for item in candidates:
        item = clip(item, max_len)
        if item and item not in result:
            result.append(item)
        if len(result) >= count:
            break
    while len(result) < count:
        result.append("围绕论文主题提炼关键依据")
    return result


def caption_lines(text: str) -> list[str]:
    captions = []
    for line in text.splitlines():
        item = strip_markdown(line)
        if re.match(r"^(图|表|Figure|Fig\.|Table)\s*\d+", item, flags=re.I):
            captions.append(clip(item, 60))
    return captions[:6]


def build_slide_plan(text: str, args: argparse.Namespace) -> dict[str, Any]:
    text = clean_source_text(text)
    title = args.title or extract_title(text)
    abstract = extract_named_block(text, ["摘要", "Abstract"]) or clip(" ".join(split_sentences(text)[:3]), 180)
    keywords = extract_named_block(text, ["关键词", "Keywords"], limit=120)
    sections = parse_sections(text)
    content_sections = [section for section in sections if normalize_space(section["title"]) != normalize_space(title)]
    if content_sections:
        sections = content_sections
    classified = classify_sections(sections, text)
    figures = caption_lines(text)

    desired_count = args.slide_count
    if desired_count is None:
        desired_count = 14
    # A complete SCUT paper report needs a cover, contents, six section
    # dividers, five content slides, and closing. Shorter requests would break
    # the chapter structure, so keep 14 as the minimum structural deck.
    desired_count = max(14, min(15, desired_count))

    focus = args.focus or "论文核心发现"
    subtitle = clip(args.focus or keywords or "论文研究汇报", 18)

    def evidence(category: str) -> str:
        return clip(classified.get(category, text), 120)

    contribution_bullets = bullets_from(classified["problem"], 3, 42)
    method_bullets = bullets_from(classified["method"], 4, 52)
    experiment_bullets = bullets_from(classified["experiment"], 3, 52)
    result_bullets = bullets_from(classified["result"], 4, 52)
    discussion_bullets = bullets_from(classified["discussion"], 3, 52)
    conclusion_bullets = bullets_from(classified["conclusion"], 3, 46)

    background_bullets = bullets_from(classified["background"], 3, 42)
    chapter_titles = ["研究背景", "问题与贡献", "方法框架", "实验设计", "结果讨论", "结论展望"]
    chapter_details = [
        clip("；".join(background_bullets[:2]), 24),
        clip("；".join(contribution_bullets[:2]), 24),
        clip("；".join(method_bullets[:2]), 24),
        clip("；".join(experiment_bullets[:2]), 24),
        clip("；".join(result_bullets[:2]), 24),
        clip("；".join(conclusion_bullets[:2]), 24),
    ]

    slides: list[dict[str, Any]] = [
        {
            "slide_type": "cover",
            "template_slide": 7,
            "layout_ref": "slide-07-cover",
            "title": title,
            "body": [],
            "source_evidence": title,
            "visual_notes": "Cover with SCUT blue visual identity, speaker, and date.",
        },
        {
            "slide_type": "toc",
            "template_slide": 8,
            "layout_ref": "slide-08-directory",
            "title": "目录",
            "body": chapter_titles,
            "details": chapter_details,
            "source_evidence": abstract,
            "visual_notes": "Six-part contents page.",
        },
        {
            "slide_type": "section",
            "template_slide": 11,
            "section_no": 1,
            "layout_ref": "slide-11-section-divider",
            "title": "研究背景",
            "subtitle": subtitle,
            "body": background_bullets[:2],
            "source_evidence": evidence("background"),
            "visual_notes": "Chapter 1 section divider.",
        },
        {
            "slide_type": "overview",
            "template_slide": 15,
            "layout_ref": "slide-15-overview-cards",
            "title": "研究背景",
            "subtitle": subtitle,
            "overview_title": "研究背景与动因",
            "overview": clip(abstract, 120),
            "body": background_bullets,
            "cards": [
                {"title": "现实背景", "body": background_bullets[0]},
                {"title": "研究动因", "body": background_bullets[1]},
                {"title": "切入问题", "body": contribution_bullets[0]},
            ],
            "source_evidence": evidence("background"),
            "visual_notes": "Background overview with three cards.",
        },
        {
            "slide_type": "section",
            "template_slide": 12,
            "section_no": 2,
            "layout_ref": "slide-12-section-divider",
            "title": "问题与贡献",
            "subtitle": focus,
            "body": contribution_bullets[:2],
            "source_evidence": evidence("problem"),
            "visual_notes": "Chapter 2 section divider.",
        },
        {
            "slide_type": "overview",
            "template_slide": 15,
            "layout_ref": "slide-15-overview-cards",
            "title": "问题与贡献",
            "subtitle": subtitle,
            "overview_title": "核心研究问题",
            "overview": clip(abstract, 120),
            "body": contribution_bullets,
            "cards": [
                {"title": "研究问题", "body": contribution_bullets[0]},
                {"title": "创新贡献", "body": contribution_bullets[1]},
                {"title": "应用价值", "body": contribution_bullets[2]},
            ],
            "source_evidence": evidence("problem"),
            "visual_notes": "Main overview with three contribution cards.",
        },
        {
            "slide_type": "section",
            "template_slide": 13,
            "section_no": 3,
            "layout_ref": "slide-13-section-divider",
            "title": "方法框架",
            "subtitle": subtitle,
            "body": method_bullets[:2],
            "source_evidence": evidence("method"),
            "visual_notes": "Chapter 3 section divider.",
        },
        {
            "slide_type": "modules",
            "template_slide": 16,
            "layout_ref": "slide-16-four-modules",
            "title": "方法框架",
            "subtitle": subtitle,
            "body": method_bullets,
            "cards": [
                {"title": "输入与数据", "body": method_bullets[0]},
                {"title": "模型或方法", "body": method_bullets[1]},
                {"title": "关键流程", "body": method_bullets[2]},
                {"title": "验证方式", "body": method_bullets[3]},
            ],
            "source_evidence": evidence("method"),
            "visual_notes": "Four horizontal method modules.",
        },
        {
            "slide_type": "section",
            "template_slide": 14,
            "section_no": 4,
            "layout_ref": "slide-14-section-divider",
            "title": "实验设计",
            "subtitle": subtitle,
            "body": experiment_bullets[:2],
            "source_evidence": evidence("experiment"),
            "visual_notes": "Chapter 4 section divider.",
        },
        {
            "slide_type": "cards",
            "template_slide": 17,
            "layout_ref": "slide-17-three-cards",
            "title": "实验与证据",
            "subtitle": subtitle,
            "body": experiment_bullets,
            "cards": [
                {"title": "数据来源", "body": experiment_bullets[0]},
                {"title": "实验设置", "body": experiment_bullets[1]},
                {"title": "评价指标", "body": experiment_bullets[2]},
            ],
            "source_evidence": evidence("experiment"),
            "visual_notes": "Three visual cards; use supplied figures when available.",
        },
        {
            "slide_type": "section",
            "template_slide": 11,
            "section_no": 5,
            "layout_ref": "slide-11-section-divider",
            "title": "结果讨论",
            "subtitle": subtitle,
            "body": result_bullets[:2],
            "source_evidence": evidence("result"),
            "visual_notes": "Chapter 5 section divider.",
        },
        {
            "slide_type": "summary",
            "template_slide": 18,
            "layout_ref": "slide-18-image-summary",
            "title": "结果与讨论",
            "subtitle": subtitle,
            "body": result_bullets,
            "overview_title": "主要发现",
            "overview": "\n".join(result_bullets[:3]),
            "note": clip("；".join(figures or discussion_bullets), 160),
            "source_evidence": evidence("result"),
            "visual_notes": "Image plus result summary page.",
        },
        {
            "slide_type": "section",
            "template_slide": 14,
            "section_no": 6,
            "layout_ref": "slide-14-section-divider",
            "title": "结论展望",
            "subtitle": focus,
            "body": conclusion_bullets[:2],
            "source_evidence": evidence("conclusion"),
            "visual_notes": "Chapter 6 section divider.",
        },
        {
            "slide_type": "thanks",
            "template_slide": 19,
            "layout_ref": "slide-19-closing",
            "title": "欢迎批评指正",
            "body": [],
            "source_evidence": "",
            "visual_notes": "Closing slide.",
        },
    ]

    if desired_count >= 15:
        slides.insert(
            2,
            {
                "slide_type": "roadmap",
                "template_slide": 10,
                "layout_ref": "slide-10-roadmap",
                "title": "研究汇报脉络",
                "body": chapter_titles[:4],
                "details": chapter_details[:4],
                "source_evidence": abstract,
                "visual_notes": "Optional roadmap page.",
            },
        )
    for idx, slide in enumerate(slides, 1):
        slide["number"] = idx

    return {
        "deck_meta": {
            "title": title,
            "speaker": args.speaker or "汇报人",
            "date": args.date or default_date(),
            "language": "zh",
            "style": "SCUT blue",
            "focus": focus,
            "slide_count": len(slides),
        },
        "slides": slides,
    }


def default_date() -> str:
    today = dt.date.today()
    return f"{today.year}年{today.month}月{today.day}日"


def iter_shapes(shapes: Any):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from iter_shapes(shape.shapes)


def text_shapes(slide: Any) -> list[Any]:
    return [shape for shape in iter_shapes(slide.shapes) if getattr(shape, "has_text_frame", False)]


def shape_text(shape: Any) -> str:
    return getattr(shape, "text", "") or ""


def first_run_style(shape: Any) -> dict[str, Any]:
    style: dict[str, Any] = {}
    if not getattr(shape, "has_text_frame", False):
        return style
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            font = run.font
            style["name"] = font.name
            style["size"] = font.size
            style["bold"] = font.bold
            style["italic"] = font.italic
            style["underline"] = font.underline
            try:
                style["rgb"] = font.color.rgb
            except Exception:
                style["rgb"] = None
            return style
    return style


def apply_run_style(shape: Any, style: dict[str, Any], font_size: int | None, font_color: tuple[int, int, int] | None) -> None:
    if not getattr(shape, "has_text_frame", False):
        return
    rgb = RGBColor(*font_color) if font_color is not None else style.get("rgb")
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            font = run.font
            if style.get("name"):
                font.name = style["name"]
            font.size = Pt(font_size) if font_size is not None else style.get("size")
            if style.get("bold") is not None:
                font.bold = style["bold"]
            if style.get("italic") is not None:
                font.italic = style["italic"]
            if style.get("underline") is not None:
                font.underline = style["underline"]
            if rgb is not None:
                font.color.rgb = rgb


def set_text(
    shape: Any,
    text: str,
    max_chars: int | None = None,
    font_size: int | None = None,
    font_color: tuple[int, int, int] | None = None,
) -> None:
    if max_chars is not None:
        text = clip(text, max_chars)
    style = first_run_style(shape)
    word_wrap = shape.text_frame.word_wrap if getattr(shape, "has_text_frame", False) else None
    auto_size = shape.text_frame.auto_size if getattr(shape, "has_text_frame", False) else None
    shape.text = text
    if getattr(shape, "has_text_frame", False):
        shape.text_frame.word_wrap = word_wrap
        shape.text_frame.auto_size = auto_size
        apply_run_style(shape, style, font_size, font_color)


def replace_contains(
    slide: Any,
    replacements: list[tuple[str, str]],
    max_chars: int | None = None,
    font_size: int | None = None,
    font_color: tuple[int, int, int] | None = None,
) -> None:
    for shape in text_shapes(slide):
        old = shape_text(shape)
        for needle, value in replacements:
            if needle in old:
                set_text(shape, value, max_chars=max_chars, font_size=font_size, font_color=font_color)
                break


def replace_sequence(
    slide: Any,
    predicate,
    values: list[str],
    max_chars: int | None = None,
    font_size: int | None = None,
    font_color: tuple[int, int, int] | None = None,
) -> None:
    queue = list(values)
    for shape in text_shapes(slide):
        if not queue:
            return
        if predicate(shape_text(shape).strip()):
            set_text(shape, queue.pop(0), max_chars=max_chars, font_size=font_size, font_color=font_color)


def cards(slide_spec: dict[str, Any], count: int) -> list[dict[str, str]]:
    data = slide_spec.get("cards") or []
    result = [{"title": item.get("title", ""), "body": item.get("body", "")} for item in data[:count]]
    while len(result) < count:
        result.append({"title": "关键要点", "body": "围绕论文主题提炼关键依据"})
    return result


def wrap_cover_title(text: str, limit: int = 24) -> str:
    text = clip(text, limit)
    if len(text) <= 12:
        return text
    midpoint = len(text) // 2
    if " " in text:
        spaces = [i for i, char in enumerate(text) if char == " "]
        split = min(spaces, key=lambda i: abs(i - midpoint))
        return text[:split].strip() + "\n" + text[split:].strip()
    split = max(8, min(12, midpoint))
    return text[:split].rstrip("，,；;。.") + "\n" + text[split:].lstrip("，,；;。.")


def apply_cover(slide: Any, spec: dict[str, Any], meta: dict[str, Any]) -> None:
    for shape in text_shapes(slide):
        old = shape_text(shape)
        if "华工蓝" in old:
            set_text(shape, wrap_cover_title(spec["title"]), font_size=24, font_color=WHITE)
        elif "South China" in old:
            set_text(shape, "South China University Of Technology", font_color=WHITE)
        elif "汇报人" in old:
            set_text(shape, f"汇报人：{compact_speaker(meta['speaker'])}", max_chars=8, font_color=WHITE)
        elif "日期" in old:
            set_text(shape, f"日期：{compact_date(meta['date'])}", max_chars=14, font_color=WHITE)


def compact_date(value: str) -> str:
    match = re.match(r"^(\d{4})年(\d{1,2})月(\d{1,2})日$", value)
    if match:
        year, month, day = match.groups()
        return f"{year}.{int(month):02d}.{int(day):02d}"
    return value


def compact_speaker(value: str) -> str:
    value = normalize_space(value)
    return value[:3] if len(value) > 3 else value


def apply_toc(slide: Any, spec: dict[str, Any]) -> None:
    topics = [clip(item, 12) for item in spec.get("body", [])[:6]]
    details = [clip(item, 22) for item in spec.get("details", [])[:6]]
    while len(topics) < 6:
        topics.append("研究内容")
    while len(details) < 6:
        details.append("论文要点概览")

    replace_contains(slide, [("在此输入摘要", clip(spec.get("source_evidence", ""), 88))], max_chars=92, font_color=WHITE)
    remove_toc_entry_groups(slide)
    add_toc_six_grid(slide, topics, details)


def remove_toc_entry_groups(slide: Any) -> None:
    for shape in list(slide.shapes):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP and shape.top > 3300000:
            shape._element.getparent().remove(shape._element)


def add_toc_six_grid(slide: Any, topics: list[str], details: list[str]) -> None:
    positions = [
        (960000, 3760000),
        (4680000, 3760000),
        (8400000, 3760000),
        (960000, 5160000),
        (4680000, 5160000),
        (8400000, 5160000),
    ]
    for idx, (topic, detail, (left, top)) in enumerate(zip(topics, details, positions), 1):
        add_toc_textbox(slide, left, top, 540000, 330000, f"{idx:02d}", 22, SCUT_BLUE, True)
        add_toc_textbox(slide, left + 610000, top + 35000, 2300000, 360000, topic, 19, SCUT_BLUE, True)
        add_toc_textbox(slide, left + 610000, top + 450000, 2550000, 520000, detail, 10.5, BLACK, False)


def add_toc_textbox(
    slide: Any,
    left: int,
    top: int,
    width: int,
    height: int,
    text: str,
    font_size: float,
    font_color: tuple[int, int, int],
    bold: bool,
) -> None:
    shape = slide.shapes.add_textbox(left, top, width, height)
    shape.text = text
    shape.text_frame.word_wrap = True
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.name = "思源黑体 Heavy" if bold else "思源黑体 CN"
            run.font.size = Pt(font_size)
            run.font.bold = bold
            run.font.color.rgb = RGBColor(*font_color)


def apply_section(slide: Any, spec: dict[str, Any]) -> None:
    body = "\n".join(f"{item}" for item in spec.get("body", [])[:2])
    section_color = WHITE if spec.get("template_slide") in {11, 12} else None
    section_no = spec.get("section_no") or max(1, spec["number"] - 2)
    replace_sequence(
        slide,
        lambda t: re.fullmatch(r"0[1-9]|[1-9]", t or "") is not None,
        [f"{section_no:02d}"],
        font_color=section_color,
    )
    replace_sequence(
        slide,
        lambda t: t in {"博学", "慎思", "明辨", "笃行"},
        [spec["title"]],
        max_chars=18,
        font_color=section_color,
    )
    replace_sequence(slide, lambda t: t in SLOGANS, [spec.get("subtitle", "")], max_chars=28, font_color=section_color)
    replace_sequence(
        slide,
        lambda t: any(fragment in t for fragment in ["云山苍苍", "先贤创业", "金银岛畔", "木棉花红"]),
        [body],
        max_chars=78,
        font_color=section_color,
    )
    replace_contains(slide, [("PRESENTATION SLIDE", "SCUT 2026\nPAPER REPORT")], max_chars=32, font_color=section_color)


def apply_roadmap(slide: Any, spec: dict[str, Any]) -> None:
    topics = [clip(item, 10) for item in spec.get("body", [])[:4]]
    details = [clip(item, 28) for item in spec.get("details", [])[:4]]
    while len(topics) < 4:
        topics.append("研究内容")
    while len(details) < 4:
        details.append("论文要点")
    replace_contains(slide, [("目录", spec["title"]), ("2026", "SCUT")], max_chars=20)
    replace_sequence(slide, lambda t: t in {"博学", "慎思", "明辨", "笃行"}, topics, max_chars=12)
    replace_sequence(slide, lambda t: t in SLOGANS, details, max_chars=30)


def apply_overview(slide: Any, spec: dict[str, Any]) -> None:
    cs = cards(spec, 3)
    replace_sequence(slide, lambda t: t in {"博学慎思", "明辨笃行"}, [spec["title"]], max_chars=18)
    replace_sequence(slide, lambda t: t in SLOGANS, [spec.get("subtitle", "")], max_chars=30)
    replace_sequence(
        slide,
        lambda t: t == "华南理工大学",
        [spec.get("overview_title", "核心研究问题")],
        max_chars=22,
        font_color=WHITE,
    )
    replace_sequence(
        slide,
        lambda t: "直属教育部" in t or "最早可溯源" in t,
        [spec.get("overview", "")],
        max_chars=115,
        font_color=WHITE,
    )
    replace_sequence(
        slide,
        lambda t: t in {"科技成果转化", "科研创新平台", "国际合作网络"},
        [item["title"] for item in cs],
        max_chars=16,
    )
    replace_sequence(
        slide,
        lambda t: any(fragment in t for fragment in ["发明专利", "科研平台", "教育合作"]),
        [item["body"] for item in cs],
        max_chars=80,
    )


def apply_modules(slide: Any, spec: dict[str, Any]) -> None:
    cs = cards(spec, 4)
    replace_sequence(slide, lambda t: t in {"博学慎思", "明辨笃行"}, [spec["title"]], max_chars=18)
    replace_sequence(slide, lambda t: t in SLOGANS, [spec.get("subtitle", "")], max_chars=30)
    module_titles = [
        ("核心攻坚 引领前沿", cs[0]["title"], None),
        ("“四链”融合 华工模式", cs[1]["title"], None),
        ("战略人才 迈向卓越", cs[2]["title"], None),
        ("立德树人 走在前列", cs[3]["title"], WHITE),
    ]
    for placeholder, value, color in module_titles:
        replace_sequence(slide, lambda t, p=placeholder: t == p, [value], max_chars=18, font_color=color)
    module_bodies = [
        ("科研经费", cs[0]["body"], None),
        ("融入发展", cs[1]["body"], None),
        ("卓越工程师", cs[2]["body"], None),
        ("低空技术", cs[3]["body"], WHITE),
    ]
    for fragment, value, color in module_bodies:
        replace_sequence(slide, lambda t, f=fragment: f in t, [value], max_chars=100, font_color=color)


def apply_cards(slide: Any, spec: dict[str, Any], image_paths: list[Path]) -> None:
    cs = cards(spec, 3)
    replace_sequence(slide, lambda t: t in {"博学慎思", "明辨笃行"}, [spec["title"]], max_chars=18)
    replace_sequence(slide, lambda t: t in SLOGANS, [spec.get("subtitle", "")], max_chars=30)
    replace_sequence(
        slide,
        lambda t: t in {"校区布局特色", "师资力量雄厚", "育人成效显著"},
        [item["title"] for item in cs],
        max_chars=14,
    )
    replace_sequence(
        slide,
        lambda t: any(fragment in t for fragment in ["一校三区", "斯坦福", "工程师的摇篮"]),
        [item["body"] for item in cs],
        max_chars=78,
    )
    replace_contains(slide, [("博学慎思  明辨笃行", "论文研究汇报")], max_chars=16, font_color=WHITE)
    keep_footer_numbers_on_one_line(slide)
    place_images(slide, image_paths, limit=3)


def apply_summary(slide: Any, spec: dict[str, Any], image_paths: list[Path]) -> None:
    replace_sequence(slide, lambda t: t in {"博学慎思", "明辨笃行"}, [spec["title"]], max_chars=18)
    replace_sequence(slide, lambda t: t in SLOGANS, [spec.get("subtitle", "")], max_chars=30)
    replace_sequence(
        slide,
        lambda t: t == "华南理工大学学生记者团",
        [spec.get("overview_title", "主要发现")],
        max_chars=20,
        font_color=WHITE,
    )
    replace_sequence(
        slide,
        lambda t: "用文字记录" in t,
        ["\n".join(spec.get("body", [])[:3])],
        max_chars=82,
        font_color=WHITE,
    )
    replace_sequence(
        slide,
        lambda t: "学生记者团多次获得" in t,
        [spec.get("note") or "结果表明，论文方法在关键指标上具有较好表现，并为后续研究提供参考。"],
        max_chars=88,
    )
    place_images(slide, image_paths, limit=1)


def keep_footer_numbers_on_one_line(slide: Any) -> None:
    for shape in text_shapes(slide):
        if shape.top < 5000000:
            continue
        if re.fullmatch(r"0[1-9]", shape_text(shape).strip()):
            shape.text_frame.word_wrap = False


def apply_thanks(slide: Any, spec: dict[str, Any], meta: dict[str, Any]) -> None:
    for shape in text_shapes(slide):
        old = shape_text(shape)
        if "欢迎批评指正" in old:
            set_text(shape, spec.get("title", "欢迎批评指正"), max_chars=18)
        elif "South China" in old:
            set_text(shape, "South China University Of Technology")
        elif "汇报人" in old:
            set_text(shape, f"汇报人：{meta['speaker']}", max_chars=14)


def sanitize_placeholders(slide: Any, meta: dict[str, Any]) -> None:
    for shape in text_shapes(slide):
        text = shape_text(shape)
        if not text:
            continue
        new_text = text.replace("鲤工仔", meta["speaker"]).replace("XX年XX月", meta["date"])
        new_text = new_text.replace("Your Title Here SCUT", "")
        new_text = new_text.replace("华工蓝幻灯片模板", meta["title"])
        if new_text != text:
            set_text(shape, new_text)


def place_images(slide: Any, image_paths: list[Path], limit: int) -> None:
    if not image_paths:
        return
    pictures = [
        shape
        for shape in iter_shapes(slide.shapes)
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and not (shape.top < 700000 and shape.left > 8000000)
    ]
    pictures.sort(key=lambda pic: pic.width * pic.height, reverse=True)
    for picture, image_path in zip(pictures[:limit], image_paths[:limit]):
        left, top, width, height = picture.left, picture.top, picture.width, picture.height
        parent = picture._element.getparent()
        parent.remove(picture._element)
        added = slide.shapes.add_picture(str(image_path), left, top, width=width, height=height)
        crop_picture_to_fill(added, image_path, width, height)


def crop_picture_to_fill(picture: Any, image_path: Path, box_width: int, box_height: int) -> None:
    try:
        from PIL import Image

        with Image.open(image_path) as image:
            image_width, image_height = image.size
    except Exception:
        return

    if image_width <= 0 or image_height <= 0 or box_width <= 0 or box_height <= 0:
        return

    picture.crop_left = 0
    picture.crop_right = 0
    picture.crop_top = 0
    picture.crop_bottom = 0

    image_ratio = image_width / image_height
    box_ratio = box_width / box_height
    if image_ratio > box_ratio:
        visible_width = image_height * box_ratio
        crop = max(0.0, min(0.49, (image_width - visible_width) / image_width / 2))
        picture.crop_left = crop
        picture.crop_right = crop
    elif image_ratio < box_ratio:
        visible_height = image_width / box_ratio
        crop = max(0.0, min(0.49, (image_height - visible_height) / image_height / 2))
        picture.crop_top = crop
        picture.crop_bottom = crop


def clone_template_slide(prs: Presentation, template_slide: int) -> Any:
    source = prs.slides[template_slide - 1]
    dest = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in list(dest.shapes):
        shape._element.getparent().remove(shape._element)

    copied_elements = []
    for shape in source.shapes:
        element = deepcopy(shape._element)
        copied_elements.append(element)
        dest.shapes._spTree.insert_element_before(element, "p:extLst")

    rid_map: dict[str, str] = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith("/notesSlide") or rel.reltype.endswith("/slideLayout"):
            continue
        target = rel.target_ref if rel.is_external else rel.target_part
        rid_map[rel.rId] = dest.part.relate_to(target, rel.reltype, rel.is_external)

    for element in copied_elements:
        replace_relationship_ids(element, rid_map)

    return dest


def replace_relationship_ids(element: Any, rid_map: dict[str, str]) -> None:
    if not rid_map:
        return
    for child in element.iter():
        for attr, value in list(child.attrib.items()):
            if value in rid_map:
                child.set(attr, rid_map[value])


def clone_template_sequence(prs: Presentation, template_slides: list[int]) -> None:
    sld_id_lst = prs.slides._sldIdLst
    clone_ids = []
    for template_slide in template_slides:
        clone_template_slide(prs, template_slide)
        clone_ids.append(list(sld_id_lst)[-1])

    for sld_id in list(sld_id_lst):
        if sld_id in clone_ids:
            continue
        rel_id = getattr(sld_id, "rId", None) or sld_id.get(f"{{{OFFICE_REL_NS}}}id")
        sld_id_lst.remove(sld_id)
        if rel_id:
            try:
                prs.part.drop_rel(rel_id)
            except KeyError:
                pass


def image_files(image_dir: Path | None) -> list[Path]:
    if not image_dir:
        return []
    exts = {".png", ".jpg", ".jpeg"}
    return sorted(path for path in image_dir.iterdir() if path.suffix.lower() in exts)


def apply_plan(plan: dict[str, Any], template: Path, output: Path, image_dir: Path | None = None) -> None:
    prs = Presentation(str(template))
    slides_spec = plan["slides"]
    clone_template_sequence(prs, [slide["template_slide"] for slide in slides_spec])
    images = image_files(image_dir)
    image_cursor = 0

    for slide, spec in zip(prs.slides, slides_spec):
        meta = plan["deck_meta"]
        slide_images: list[Path] = []
        if spec["slide_type"] in {"cards", "summary"} and images:
            take = 3 if spec["slide_type"] == "cards" else 1
            slide_images = images[image_cursor : image_cursor + take]
            image_cursor += len(slide_images)

        if spec["slide_type"] == "cover":
            apply_cover(slide, spec, meta)
        elif spec["slide_type"] == "toc":
            apply_toc(slide, spec)
        elif spec["slide_type"] == "roadmap":
            apply_roadmap(slide, spec)
        elif spec["slide_type"] == "section":
            apply_section(slide, spec)
        elif spec["slide_type"] == "overview":
            apply_overview(slide, spec)
        elif spec["slide_type"] == "modules":
            apply_modules(slide, spec)
        elif spec["slide_type"] == "cards":
            apply_cards(slide, spec, slide_images)
        elif spec["slide_type"] == "summary":
            apply_summary(slide, spec, slide_images)
        elif spec["slide_type"] == "thanks":
            apply_thanks(slide, spec, meta)
        sanitize_placeholders(slide, meta)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output))


def load_plan(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def write_plan(plan: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Generate an editable SCUT blue PPTX from a paper.")
    parser.add_argument("--input", type=Path, help="Paper PDF, Markdown, or TXT file.")
    parser.add_argument("--output", type=Path, required=True, help="Output PPTX path.")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE, help="Template PPTX path.")
    parser.add_argument("--plan", type=Path, help="Existing slide_plan.json to use instead of auto-planning.")
    parser.add_argument("--plan-output", type=Path, help="Where to write the generated plan JSON.")
    parser.add_argument("--speaker", default="汇报人", help="Speaker name shown on cover and closing slides.")
    parser.add_argument("--date", help="Cover date, e.g. 2026年5月18日.")
    parser.add_argument("--title", help="Override deck title.")
    parser.add_argument("--slide-count", type=int, help="Desired slide count. Auto-planning keeps 14 as the minimum for six chapter dividers.")
    parser.add_argument("--focus", help="Presentation focus, e.g. 方法与实验结果.")
    parser.add_argument("--image-dir", type=Path, help="Optional folder of figures to place into image layouts.")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.plan:
        plan = load_plan(args.plan)
    else:
        if not args.input:
            raise SystemExit("--input is required unless --plan is provided.")
        plan = build_slide_plan(read_text(args.input), args)

    plan_path = args.plan_output or args.output.with_suffix(".slide_plan.json")
    write_plan(plan, plan_path)
    apply_plan(plan, args.template, args.output, args.image_dir)
    print(f"Wrote {args.output}")
    print(f"Wrote {plan_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
